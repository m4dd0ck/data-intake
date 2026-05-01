"""Read client exports as text tables.

Everything is loaded as text on purpose: the audit's job is to find out what types the data
really has, so no reader is allowed to guess and silently coerce first.
"""

import shutil
import tempfile
from collections.abc import Iterable
from datetime import date, datetime, time
from pathlib import Path

import duckdb
from openpyxl import load_workbook

from data_intake.models import RawTable

SUPPORTED = {".csv", ".txt", ".xlsx", ".xlsm"}
HEADER_SCAN_ROWS = 20
# Limits for hostile or accidental giant files; well above any real small-business export.
MAX_FILE_BYTES = 200 * 1024 * 1024
MAX_ROWS = 1_000_000
MAX_COLUMNS = 500
GLOB_CHARS = set("*?[]{}")


class UnsupportedFileError(ValueError):
    """Raised for a file type the audit cannot read."""


class FileTooLargeError(ValueError):
    """Raised when a file exceeds the size, row or column limits."""


def load_folder(folder: Path) -> list[RawTable]:
    """Load every supported file in a folder (not recursive), sorted by name.

    Raises:
        FileNotFoundError: If the folder does not exist or holds no supported files.
    """
    if not folder.is_dir():
        raise FileNotFoundError(f"Not a folder: {folder}")
    # Reason: skip symlinks so a link inside an unzipped client archive cannot pull a local
    # file (say ~/.ssh/config) into a report that gets sent back.
    files = sorted(
        p
        for p in folder.iterdir()
        if p.suffix.lower() in SUPPORTED and p.is_file() and not p.is_symlink()
    )
    if not files:
        raise FileNotFoundError(f"No CSV or Excel files in {folder}")
    tables: list[RawTable] = []
    for path in files:
        tables.extend(load_file(path))
    return tables


def load_file(path: Path) -> list[RawTable]:
    """One table per CSV, one per non-empty Excel sheet."""
    if path.stat().st_size > MAX_FILE_BYTES:
        raise FileTooLargeError(f"{path.name} is over {MAX_FILE_BYTES // 2**20} MB")
    suffix = path.suffix.lower()
    if suffix in {".csv", ".txt"}:
        return [_build_table(path.stem, path, _read_csv(path))]
    if suffix in {".xlsx", ".xlsm"}:
        return _read_workbook(path)
    raise UnsupportedFileError(f"Unsupported file type: {path.name}")


def _read_csv(path: Path) -> list[list[str | None]]:
    # DuckDB treats the path as a glob, so a file literally named "Sales [2024].csv" is read
    # through a copy with a plain name.
    if GLOB_CHARS & set(path.name):
        with tempfile.TemporaryDirectory() as tmp:
            plain = Path(tmp) / "export.csv"
            shutil.copyfile(path, plain)
            return _read_csv(plain)
    # Reason: header=false so title rows and odd headers reach our own header detection.
    with duckdb.connect() as connection:
        relation = connection.execute(
            "select * from read_csv(?, all_varchar = true, header = false, sample_size = -1)",
            [str(path)],
        )
        rows = [list(row) for row in relation.fetchmany(MAX_ROWS + 1)]
    _check_shape(path, len(rows), max((len(r) for r in rows), default=0))
    return rows


def _read_workbook(path: Path) -> list[RawTable]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    tables = []
    try:
        for sheet in workbook.worksheets:
            # Reason: read-only mode trusts the sheet's declared size and pads to it; a sheet
            # claiming A1:XFD1048576 would build billions of empty cells.
            sheet.reset_dimensions()
            rows = []
            for row in sheet.iter_rows(values_only=True):
                rows.append([_cell_text(v) for v in row[: MAX_COLUMNS + 1]])
                if len(rows) > MAX_ROWS:
                    break
            _check_shape(path, len(rows), max((len(r) for r in rows), default=0))
            if any(any(cell is not None for cell in row) for row in rows):
                name = path.stem if len(workbook.worksheets) == 1 else f"{path.stem}/{sheet.title}"
                tables.append(_build_table(name, path, rows))
    finally:
        workbook.close()
    return tables


def _check_shape(path: Path, rows: int, columns: int) -> None:
    if rows > MAX_ROWS or columns > MAX_COLUMNS:
        raise FileTooLargeError(
            f"{path.name} exceeds {MAX_ROWS:,} rows or {MAX_COLUMNS} columns; split it first"
        )


def _cell_text(value: object) -> str | None:
    """Excel cell value as the text a person would see, without reformatting numbers."""
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date().isoformat() if value.time() == time(0) else value.isoformat(" ")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    text = str(value)
    return text if text != "" else None


def _build_table(name: str, source: Path, rows: list[list[str | None]]) -> RawTable:
    rows = [row for row in rows if any(cell not in (None, "") for cell in row)]
    header_index = find_header_row(rows)
    width = max((len(row) for row in rows), default=0)
    header = rows[header_index] if rows else []
    columns, renamed = clean_headers(list(header) + [None] * (width - len(header)))
    body = [list(row) + [None] * (width - len(row)) for row in rows[header_index + 1 :]]
    return RawTable(name, source, columns, body, header_index, renamed)


def find_header_row(rows: list[list[str | None]]) -> int:
    """First row that is filled across most of the table's width.

    Exports often start with a title and a "generated on" line; those rows fill one or two
    cells, while the header fills most columns.
    """
    sample = rows[:HEADER_SCAN_ROWS]
    width = max((_filled(row) for row in sample), default=0)
    for index, row in enumerate(sample):
        if _filled(row) >= max(2, width / 2):
            return index
    return 0


def _filled(row: Iterable[str | None]) -> int:
    return sum(1 for cell in row if cell not in (None, ""))


def clean_headers(header: list[str | None]) -> tuple[list[str], dict[str, str]]:
    """Trim names, name blanks ``unnamed_N`` and suffix repeats; report what changed."""
    columns: list[str] = []
    renamed: dict[str, str] = {}
    seen: dict[str, int] = {}
    for position, raw in enumerate(header, start=1):
        original = "" if raw is None else str(raw)
        name = original.strip() or f"unnamed_{position}"
        if name in seen:
            seen[name] += 1
            name = f"{name}_{seen[name]}"
        else:
            seen[name] = 1
        if name != original:
            renamed[name] = original
        columns.append(name)
    return columns, renamed
