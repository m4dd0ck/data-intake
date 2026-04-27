"""Messy demo exports for two small businesses, with every planted problem listed.

``PLANTED`` is the answer key: the recall test asserts the audit finds each one, and
``CLEAN_COLUMNS`` lists columns that must come back with no findings at all.
"""

import csv
import random
from dataclasses import dataclass
from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook

AS_OF = date(2024, 12, 31)


@dataclass(frozen=True)
class Planted:
    business: str
    table: str
    code: str
    column: str | None


PLANTED = [
    Planted("shop", "orders", "mixed_date_formats", "order_date"),
    Planted("shop", "orders", "numbers_as_text", "amount"),
    Planted("shop", "orders", "duplicate_rows", None),
    Planted("shop", "orders", "duplicate_keys", "order_id"),
    Planted("shop", "orders", "blank_placeholders", "customer_id"),
    Planted("shop", "orders", "negative_amounts", "amount"),
    Planted("shop", "orders", "outliers", "amount"),
    Planted("shop", "orders", "future_dates", "order_date"),
    Planted("shop", "orders", "inconsistent_labels", "country"),
    Planted("shop", "orders", "orphan_keys", "customer_id"),
    Planted("shop", "customers", "title_rows", None),
    Planted("shop", "customers", "bad_headers", None),
    Planted("shop", "customers", "empty_columns", None),
    Planted("shop", "customers", "near_duplicate_records", "name"),
    Planted("shop", "customers", "near_duplicate_records", "email"),
    Planted("shop", "products", "mixed_currencies", "price"),
    Planted("studio", "invoices/2024", "excel_serial_dates", "issued"),
    Planted("studio", "timesheets", "orphan_keys", "client_id"),
    Planted("studio", "timesheets", "outliers", "hours"),
    Planted("studio", "clients", "stray_whitespace", "city"),
]
CLEAN_COLUMNS = [
    ("shop", "products", "cost"),
    ("shop", "orders", "status"),
    ("studio", "invoices/2023", "amount"),
    ("studio", "timesheets", "person"),
]

FIRST = ["Ana", "Ben", "Chloe", "Dev", "Ema", "Finn", "Grace", "Hugo", "Isla", "Jon", "Kai", "Lena"]
LAST = ["Harbor", "Pine", "Stone", "Maple", "Rivers", "Clark", "Webb", "Ortiz", "Nakamura", "Price"]
COUNTRIES = ["US", "US", "US", "GB", "CA", "DE"]


def build_demo(out_dir: Path) -> dict[str, Path]:
    """Write both businesses' exports under ``out_dir``; return their folders."""
    rng = random.Random(7)
    shop, studio = out_dir / "shop", out_dir / "studio"
    shop.mkdir(parents=True, exist_ok=True)
    studio.mkdir(parents=True, exist_ok=True)
    _shop(rng, shop)
    _studio(rng, studio)
    return {"shop": shop, "studio": studio}


def _write_csv(path: Path, header: list[str], rows: list[list[object]]) -> None:
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(header)
        writer.writerows(rows)


def _shop(rng: random.Random, folder: Path) -> None:
    customers: list[list[str]] = []
    for number in range(1, 121):
        name = f"{rng.choice(FIRST)} {rng.choice(LAST)}"
        email = f"{name.lower().replace(' ', '.')}{number}@mail.com"
        signup = date(2023, 1, 1) + timedelta(days=rng.randint(0, 600))
        customers.append([f"C{number:04d}", name, email, rng.choice(COUNTRIES), signup.isoformat()])
    # Near-duplicates: a record re-entered with one letter missing, and an email in upper case.
    _, name, email, country, signup_text = customers[4]
    customers.append(["C0121", name[:-3] + name[-2:], email, country, signup_text])
    _, name, email, country, signup_text = customers[9]
    customers.append(["C0122", name, email.upper(), country, signup_text])

    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Customers"
    sheet.append(["Pinecrest Goods - customer export"])
    sheet.append(["Generated 2024-12-31 by StoreAdmin"])
    sheet.append(
        ["customer_id", "name", "email", "country", "signup_date", "phone", "phone", "notes"]
    )
    for row in customers:
        sheet.append(
            [*row, f"555-{rng.randint(1000, 9999)}", f"555-{rng.randint(1000, 9999)}", None]
        )
    workbook.save(folder / "customers.xlsx")

    # Reason: six customers who ordered are missing from the customer export (orphans).
    known = [c[0] for c in customers] + [f"C09{n:02d}" for n in range(6)]
    orders: list[list[object]] = []
    day = date(2024, 1, 1)
    for number in range(1, 601):
        day += timedelta(days=rng.choice([0, 0, 1]))
        amount = round(rng.lognormvariate(3.8, 0.5), 2)
        when = day.isoformat() if rng.random() > 0.15 else f"{day.month}/{day.day}/{day.year}"
        shown = f"${amount:,.2f}" if rng.random() < 0.3 else f"{amount:.2f}"
        country = rng.choice(COUNTRIES)
        if country == "US":
            country = rng.choice(["US", "US", "US", "us ", "Us"])
        orders.append([f"{1000 + number}", when, rng.choice(known), country, shown, "paid"])
    for index in rng.sample(range(600), 8):
        orders[index][2] = rng.choice(["N/A", "-"])
    orders[50][4] = "-45.00"  # a refund entered as an order
    orders[70][4] = "25000.00"  # an extra zero or two
    orders[-1][1], orders[-2][1] = "2025-01-15", "2025-02-03"  # year typos
    for index in rng.sample(range(600), 5):  # re-used ids with different amounts
        clash = list(orders[index])
        clash[4] = f"{float(str(clash[4]).replace('$', '').replace(',', '')) + 10:.2f}"
        orders.append(clash)
    orders += [list(orders[i]) for i in rng.sample(range(600), 12)]  # exact copies
    header = ["order_id", "order_date", "customer_id", "country", "amount", "status"]
    _write_csv(folder / "orders.csv", header, orders)

    products: list[list[object]] = []
    for number in range(1, 41):
        cost = round(rng.uniform(4, 60), 2)
        price = round(cost * rng.uniform(1.8, 2.6), 2)
        symbol = "€" if number % 7 == 0 else "$"
        products.append(
            [f"SKU-{number:03d}", f"Item {number}", f"{symbol}{price:.2f}", f"{cost:.2f}"]
        )
    _write_csv(folder / "products.csv", ["sku", "product_name", "price", "cost"], products)


def _studio(rng: random.Random, folder: Path) -> None:
    clients: list[list[object]] = []
    for number in range(1, 16):
        name = f"{rng.choice(LAST)} {rng.choice(['Labs', 'Studio', 'Group', 'Co'])}"
        city = rng.choice(["Leeds", "York", "Bristol", "Leeds "])
        clients.append([f"CL{number:02d}", name, city, f"20{rng.randint(19, 23)}"])
    _write_csv(folder / "clients.csv", ["client_id", "client_name", "city", "since"], clients)

    workbook = Workbook()
    first = workbook.active
    assert first is not None
    first.title = "2023"
    header = ["invoice_id", "client_id", "issued", "amount", "tax", "paid_on"]
    for sheet, year in ((first, 2023), (workbook.create_sheet("2024"), 2024)):
        sheet.append(header)
        for number in range(1, 41):
            issued = date(year, 1, 1) + timedelta(days=rng.randint(0, 330))
            paid = issued + timedelta(days=rng.randint(5, 60))
            # 2024 was exported from a sheet where the date column lost its formatting.
            shown = (issued - date(1899, 12, 30)).days if year == 2024 else issued
            sheet.append([f"INV-{year}-{number:03d}", rng.choice(clients)[0], shown,
                          round(rng.uniform(800, 6000), 2), "20%", paid])  # fmt: skip
    workbook.save(folder / "invoices.xlsx")

    people = ["Rosa", "Theo", "Mina"]
    rows: list[list[object]] = []
    for day in range(200):
        when = date(2024, 1, 1) + timedelta(days=day)
        if when.weekday() < 5:
            for person in people:
                rows.append(
                    [when.isoformat(), rng.choice(clients)[0], person, rng.choice([4, 6, 7.5, 8])]
                )
    rows[30][1] = rows[90][1] = "CL99"  # a client missing from the client list
    rows[120][3] = 60  # a week's hours typed into one day
    _write_csv(folder / "timesheets.csv", ["date", "client_id", "person", "hours"], rows)
